"""Create self-contained Word copies of the Markdown practice reports.

Run with a Python environment containing python-docx and Pillow:
    python scripts/export_reports_docx.py
"""

from __future__ import annotations

import re
from pathlib import Path
from urllib.parse import urljoin

from PIL import Image
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "reports"
OUTPUT = REPORTS / "word"
GITHUB = "https://github.com/Pondaaaaaa/GO/blob/main/"
INLINE = re.compile(r"(!?\[([^]]+)\]\(([^)]+)\)|\*\*(.+?)\*\*|`([^`]+)`)")
IMAGE_LINK = re.compile(r"(?<!!)\[([^]]+)\]\((images/[^)]+\.(?:png|jpg|jpeg))\)", re.I)


def set_font(style, name: str, size: float, bold: bool = False) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), name)


def prepare_document() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21)
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.header_distance = sec.footer_distance = Cm(0.9)

    styles = doc.styles
    set_font(styles["Normal"], "Calibri", 10.5)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.12
    set_font(styles["Title"], "Calibri", 17, True)
    styles["Title"].paragraph_format.space_after = Pt(14)
    styles["Title"].paragraph_format.keep_with_next = True
    title_properties = styles["Title"].element.get_or_add_pPr()
    for border in title_properties.findall(qn("w:pBdr")):
        title_properties.remove(border)
    for name, size in (("Heading 1", 13), ("Heading 2", 11.5), ("Heading 3", 10.5)):
        set_font(styles[name], "Calibri", size, True)
        styles[name].paragraph_format.space_before = Pt(11)
        styles[name].paragraph_format.space_after = Pt(5)
        styles[name].paragraph_format.keep_with_next = True
    for name in ("List Bullet", "List Number"):
        set_font(styles[name], "Calibri", 10.5)
        styles[name].paragraph_format.space_after = Pt(3)
    if "Code" not in styles:
        from docx.enum.style import WD_STYLE_TYPE

        styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
    set_font(styles["Code"], "Consolas", 8.3)
    styles["Code"].paragraph_format.space_after = Pt(0)
    styles["Code"].paragraph_format.line_spacing = 1.0
    styles["Code"].paragraph_format.left_indent = Cm(0.35)
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.keywords = ""
    return doc


def add_hyperlink(paragraph, label: str, target: str) -> None:
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    link_id = paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True)
    element = OxmlElement("w:hyperlink")
    element.set(qn("r:id"), link_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "195A85")
    properties.append(color)
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    element.append(run)
    paragraph._p.append(element)


def resolve_link(markdown: Path, target: str) -> str:
    if target.startswith(("https://", "http://")):
        return target
    relative = (markdown.parent / target).resolve().relative_to(ROOT)
    return urljoin(GITHUB, relative.as_posix())


def add_inline(paragraph, source: str, markdown: Path) -> None:
    cursor = 0
    for match in INLINE.finditer(source):
        if match.start() > cursor:
            paragraph.add_run(source[cursor : match.start()])
        if match.group(2) is not None:
            label, target = match.group(2).replace("`", ""), match.group(3)
            if match.group(1).startswith("!"):
                paragraph.add_run(label)
            else:
                add_hyperlink(paragraph, label, resolve_link(markdown, target))
        elif match.group(4) is not None:
            paragraph.add_run(match.group(4).replace("`", "")).bold = True
        else:
            run = paragraph.add_run(match.group(5))
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        cursor = match.end()
    if cursor < len(source):
        paragraph.add_run(source[cursor:])


def add_picture(doc: Document, path: Path, caption: str) -> None:
    if not path.is_file():
        raise FileNotFoundError(path)
    with Image.open(path) as image:
        width_px, height_px = image.size
    max_width, max_height = 17.2, 22.0
    ratio = min(max_width / width_px, max_height / height_px)
    width = width_px * ratio
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.add_run().add_picture(str(path), width=Cm(width))
    if caption:
        text = doc.add_paragraph(caption)
        text.alignment = WD_ALIGN_PARAGRAPH.CENTER
        text.paragraph_format.space_after = Pt(8)
        for run in text.runs:
            run.font.size = Pt(8.5)


def table_cells(line: str) -> list[str]:
    return [part.strip() for part in line.strip().strip("|").split("|")]


def is_separator(line: str) -> bool:
    return bool(re.fullmatch(r"\s*\|?[\s|:\-]+\|?\s*", line))


def add_table(doc: Document, lines: list[str], markdown: Path) -> None:
    rows = [table_cells(line) for line in lines if not is_separator(line)]
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for row_idx, values in enumerate(rows):
        row = table.rows[row_idx]
        for col_idx, value in enumerate(values):
            cell = row.cells[col_idx]
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            add_inline(paragraph, value, markdown)
            for run in paragraph.runs:
                run.font.size = Pt(9)
                if row_idx == 0:
                    run.bold = True
        if row_idx == 0:
            trpr = row._tr.get_or_add_trPr()
            trpr.append(OxmlElement("w:tblHeader"))
            for cell in row.cells:
                tcpr = cell._tc.get_or_add_tcPr()
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "E9EDF0")
                tcpr.append(shade)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_code_line(doc: Document, line: str, compact: bool = False) -> None:
    p = doc.add_paragraph(style="Code")
    run = p.add_run(line if line else " ")
    if compact:
        run.font.size = Pt(7.5)
        p.paragraph_format.line_spacing = 0.9


def convert(markdown: Path) -> Path:
    doc = prepare_document()
    if markdown.stem == "pz05":
        normal = doc.styles["Normal"]
        normal.font.size = Pt(10)
        normal.paragraph_format.space_after = Pt(4)
        normal.paragraph_format.line_spacing = 1.05
    lines = markdown.read_text(encoding="utf-8").splitlines()
    in_code = False
    embedded: set[Path] = set()
    index = 0
    while index < len(lines):
        line = lines[index]
        if line.startswith("```"):
            in_code = not in_code
            index += 1
            continue
        if in_code:
            add_code_line(doc, line)
            index += 1
            continue
        if not line.strip():
            index += 1
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", line)
        if heading:
            level = len(heading.group(1))
            title = heading.group(2).replace(" — ", " ") if level == 1 else heading.group(2)
            p = doc.add_paragraph(style="Title" if level == 1 else f"Heading {min(level - 1, 3)}")
            add_inline(p, title, markdown)
            index += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_table(doc, table_lines, markdown)
            continue
        image = re.fullmatch(r"!\[([^]]+)\]\(([^)]+)\)", line.strip())
        if image:
            path = (markdown.parent / image.group(2)).resolve()
            add_picture(doc, path, image.group(1))
            embedded.add(path)
            index += 1
            continue
        numbered = re.match(r"^\d+\.\s+(.+)$", line)
        bullet = re.match(r"^[-*]\s+(.+)$", line)
        content = numbered.group(1) if numbered else bullet.group(1) if bullet else line
        p = doc.add_paragraph(style="Normal" if numbered else "List Bullet" if bullet else "Normal")
        if numbered:
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.first_line_indent = Cm(-0.45)
            p.add_run(line.split(".", 1)[0] + ". ")
        add_inline(p, content, markdown)
        next_index = index + 1
        while next_index < len(lines) and not lines[next_index].strip():
            next_index += 1
        if next_index < len(lines) and lines[next_index].lstrip().startswith("!["):
            p.paragraph_format.keep_with_next = True
        for image_link in IMAGE_LINK.finditer(content):
            path = (markdown.parent / image_link.group(2)).resolve()
            if path not in embedded:
                p.paragraph_format.keep_with_next = True
                add_picture(doc, path, image_link.group(1))
                embedded.add(path)
        index += 1

    if markdown.stem == "pz08":
        log = REPORTS / "logs" / "pz08-http.txt"
        doc.add_heading("Приложение Фактический HTTP лог", 1)
        for line in log.read_text(encoding="utf-8").splitlines():
            add_code_line(doc, line, compact=True)

    OUTPUT.mkdir(parents=True, exist_ok=True)
    target = OUTPUT / f"{markdown.stem}.docx"
    doc.save(target)
    return target


def main() -> None:
    paths = sorted(REPORTS.glob("pz[0-9][0-9].md"))
    if not paths:
        raise SystemExit("No practice reports found")
    for markdown in paths:
        print(convert(markdown))


if __name__ == "__main__":
    main()
